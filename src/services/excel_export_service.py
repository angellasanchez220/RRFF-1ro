import io
import math
import copy
from pathlib import Path
from datetime import date
from sqlalchemy import text
import pandas as pd
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from api.db import engine
from api.routes.sop import _get_mes_map, _get_semanas_fact_ventas

CATEGORY_SHEET_MAP = {
    "CINTAS": "Cintas",
    "MONTAJE": "Montaje",
    "TORNILLOS": "Tornillos",
    "SELLOS": "Sellos",
    "SOBERBIO": "Soberbio",
    "GANCHOS": "Ganchos",
    "FIELTROS": "Fieltros",
    "SEGURIDAD": "Seguridad",
    "DEMARCATORIA": "Demarcatoria",
    "CLAVOS": "Clavos"
}

def _safe_float(v, default=0.0):
    if v is None:
        return default
    try:
        f = float(v)
        if math.isnan(f):
            return default
        return f
    except:
        return default

def _safe_int(v, default=0):
    return int(round(_safe_float(v, default)))

def get_sku_export_data(sku: str, familias_map: dict = None) -> dict:
    """Extrae la informaci\u00f3n consolidada de un SKU para exportaci\u00f3n a Excel."""
    with engine.connect() as conn:
        df_sop = pd.read_sql(text("SELECT * FROM planificacion_sop WHERE (sku = :s OR codigo_femaco = :s)"), conn, params={"s": sku})
        if df_sop.empty:
            return None
        row_sop = df_sop.iloc[0].to_dict()
        
        # Obtener el SKU real canónico desde la base de datos (por si se buscó por código femaco)
        real_sku = str(row_sop.get("sku", sku))
        codigo_value = row_sop.get("codigo_femaco")
        real_codigo = "" if pd.isna(codigo_value) else str(codigo_value).strip().upper()

        sem_map = _get_semanas_fact_ventas(conn)
        ventas_sku = sem_map.get(real_sku, {})

        df_tr = pd.read_sql(text("""
            SELECT codigo_envio, cantidad, COALESCE(eta_ajustada, fecha_disponibilidad_real) as eta
            FROM control_embarques
            WHERE (sku = :s OR codigo_femaco = :s) AND estado IN ('EN_TRANSITO','EN_AFORO','RETRASADO')
            ORDER BY fecha_eta ASC NULLS LAST
            LIMIT 3
        """), conn, params={"s": sku})
        pedidos = df_tr.to_dict(orient="records")

        total_transito = _safe_int(row_sop.get("cantidad_transito"))

        obs_row = conn.execute(text("SELECT observacion FROM sku_observaciones WHERE sku=:s"), {"s": real_sku}).fetchone()
        observacion = obs_row.observacion if obs_row and obs_row.observacion else ""
        
        # Familia de Reemplazo
        if familias_map and real_codigo in familias_map and len(familias_map[real_codigo]["familia_skus"]) > 1:
            fam = familias_map[real_codigo]
            comp_strs = []
            for c in fam["familia_skus"]:
                nt_str = " [NO TRANSFORMABLE]" if c.get("no_transformable") else ""
                comp_strs.append(
                    f"CÓD {c['codigo_femaco']} / SKU {c['sku']} "
                    f"Stock {int(c['stock_act'])} / RV {int(c['ritmo_mensual'])}{nt_str}"
                )
            fam_str = f"Familia SKU:\n" + "\n".join(comp_strs) + f"\n\nStock bruto familia: {int(fam['stock_bruto_familia'])} uds."
            
            reemplazos = [r['sku'] for r in fam["reemplazos_validos"]]
            if reemplazos:
                fam_str += f"\n\nReemplazos válidos para CÓD {real_codigo}:\n" + ", ".join(reemplazos)
            
            if observacion:
                observacion = observacion + "\n\n" + fam_str
            else:
                observacion = fam_str

    ue = _safe_float(row_sop.get("ump"), 1.0)
    if ue == 0:
        ue = 1.0

    inv_u = _safe_int(row_sop.get("stock_act"))
    hc_u = _safe_int(ventas_sku.get("stock_fisico_matrix"))
    
    mes_map = _get_mes_map()
    sell_out = []
    sell_in = []
    meses_nombres = []
    max_so = 0
    max_si = 0

    for mes_num, nombre, abrev, yr_db, yr_label in mes_map:
        col_so = f"sellout_{abrev}_{yr_db}"
        col_si = f"sellin_{abrev}_{yr_db}"
        so_val = _safe_int(row_sop.get(col_so))
        si_val = _safe_int(row_sop.get(col_si))
        sell_out.append(so_val)
        sell_in.append(si_val)
        meses_nombres.append(f"{nombre[:3]} {yr_label}")
        if so_val > max_so: max_so = so_val
        if si_val > max_si: max_si = si_val

    data = {
        "sku": real_sku,
        "categoria": str(row_sop.get("categoria") or "SIN CATEGORIA").strip().upper(),
        "codigo": str(row_sop.get("codigo_femaco") or ""),
        "nombre": str(row_sop.get("nombre_producto") or ""),
        "estado": str(row_sop.get("estado") or ""),
        "ue": ue,
        
        "anio_labels": meses_nombres,
        "sell_out": sell_out,
        "sell_in": sell_in,
        
        "inv_u": inv_u,
        "inv_c": round(inv_u / ue, 2),
        
        "tra_u": total_transito,
        "tra_c": round(total_transito / ue, 2),
        
        "hc_u": hc_u,
        "hc_c": round(hc_u / ue, 2),
        
        "pedidos": pedidos,
        
        "max_so": max_so,
        "max_si": max_si,
        
        "s1": _safe_int(ventas_sku.get("sem1_uds")),
        "s2": _safe_int(ventas_sku.get("sem2_uds")),
        "s3": _safe_int(ventas_sku.get("sem3_uds")),
        "s4": _safe_int(ventas_sku.get("sem4_uds")),
        "tot4s": _safe_int(ventas_sku.get("total_4_sem_verificado")),
        
        "obj": _safe_int(row_sop.get("stock_objetivo")),
        "sug_bruta": _safe_int(row_sop.get("sugerencia_bruta")),
        "compra": _safe_int(row_sop.get("sugerencia_final")),
        "cob": round(float(row_sop.get("stock_act") or 0) / _safe_float(row_sop.get("sug_ritmo_mensual")), 1) if _safe_float(row_sop.get("sug_ritmo_mensual")) > 0 else 999.0,
        "quiebre": str(row_sop.get("fecha_estimada_quiebre") or ""),
        "motivo": str(row_sop.get("explicacion_compra_dinamica") or row_sop.get("explicacion_compra") or ""),
        "obs": observacion
    }
    return data

def copiar_bloque(ws: Worksheet, fila_origen_inicio: int, fila_origen_fin: int, fila_destino_inicio: int):
    offset = fila_destino_inicio - fila_origen_inicio
    
    for row in range(fila_origen_inicio, fila_origen_fin + 1):
        if row in ws.row_dimensions:
            ws.row_dimensions[row + offset].height = ws.row_dimensions[row].height

        for col in range(1, ws.max_column + 1):
            source_cell = ws.cell(row=row, column=col)
            target_cell = ws.cell(row=row + offset, column=col)

            target_cell.value = source_cell.value
            
            if source_cell.has_style:
                target_cell.font = copy.copy(source_cell.font)
                target_cell.border = copy.copy(source_cell.border)
                target_cell.fill = copy.copy(source_cell.fill)
                target_cell.number_format = copy.copy(source_cell.number_format)
                target_cell.protection = copy.copy(source_cell.protection)
                target_cell.alignment = copy.copy(source_cell.alignment)

    merges_to_add = []
    for merge_range in list(ws.merged_cells.ranges):
        if fila_origen_inicio <= merge_range.min_row <= fila_origen_fin:
            min_col, min_row, max_col, max_row = merge_range.bounds
            new_merge = f"{openpyxl.utils.get_column_letter(min_col)}{min_row + offset}:{openpyxl.utils.get_column_letter(max_col)}{max_row + offset}"
            merges_to_add.append(new_merge)
            
    for nm in merges_to_add:
        ws.merge_cells(nm)

def poblar_bloque(ws: Worksheet, fila_inicio: int, data: dict):
    def set_val(col, row_offset, val):
        ws.cell(row=fila_inicio + row_offset, column=col, value=val)

    set_val(2, 0, data["sku"])
    set_val(3, 0, data["codigo"])
    set_val(4, 0, data["nombre"])
    set_val(7, 0, data["estado"])
    set_val(9, 0, data["ue"])
    
    set_val(2, 6, data["inv_u"])
    set_val(3, 6, data["inv_c"])
    set_val(2, 7, data["tra_u"])
    set_val(3, 7, data["tra_c"])
    set_val(2, 8, data["hc_u"])
    set_val(3, 8, data["hc_c"])
    
    set_val(10, 6, data["max_so"])
    set_val(10, 7, data["max_si"])
    set_val(10, 10, data["cob"])
    
    set_val(2, 11, data["s1"])
    set_val(3, 11, data["s2"])
    set_val(4, 11, data["s3"])
    set_val(5, 11, data["s4"])
    set_val(6, 11, data["tot4s"])
    set_val(7, 11, data["obj"])

    for i, label in enumerate(data["anio_labels"]):
        set_val(2 + i, 1, label)
        
    for i in range(12):
        set_val(2 + i, 2, data["sell_out"][i])
        set_val(2 + i, 3, data["sell_in"][i])

    for col_idx in [6, 7, 8]:
        for row_offset in [5, 6, 7, 8]:
            set_val(col_idx, row_offset, "")

    pedidos = data["pedidos"]
    for i in range(min(3, len(pedidos))):
        col_idx = 6 + i  # F, G, H
        set_val(col_idx, 5, pedidos[i]["codigo_envio"])
        set_val(col_idx, 6, pedidos[i]["cantidad"])
        eta = pedidos[i]["eta"]
        if pd.notna(eta) and eta and str(eta) != "NaT":
            try:
                set_val(col_idx, 7, pd.to_datetime(eta).strftime("%d-%m-%Y"))
            except Exception:
                set_val(col_idx, 7, str(eta)[:10])
        else:
            set_val(col_idx, 7, "-")
        set_val(col_idx, 8, "-")  # Duración no disponible

    # Sugerencia bruta
    set_val(12, 6, "Sugerencia bruta")
    sug_bruta = data.get("sug_bruta", 0)
    set_val(13, 6, f"{sug_bruta} unidades" if sug_bruta > 0 else "-")
    
    # Sugerencia final (ajustada a UMP)
    set_val(12, 7, "Sugerencia final (ajustada a UMP)")
    sug_final = data.get("compra", 0)
    set_val(13, 7, f"{sug_final} unidades" if sug_final > 0 else "-")
    
    # Cobertura
    set_val(12, 8, "Cobertura")
    cob = data.get("cob")
    if pd.notna(cob) and cob < 999.0:
        set_val(13, 8, f"{str(cob).replace('.', ',')} meses")
    else:
        set_val(13, 8, "+10 meses" if pd.notna(cob) else "-")
        
    # Quiebre
    set_val(12, 9, "Quiebre")
    quiebre = data.get("quiebre")
    if pd.notna(quiebre) and quiebre and str(quiebre) != "NaT":
        try:
            set_val(13, 9, pd.to_datetime(quiebre).strftime("%d-%m-%Y"))
        except Exception:
            set_val(13, 9, str(quiebre)[:10])
    else:
        set_val(13, 9, "-")
        
    # Motivo
    set_val(12, 10, "Motivo")
    set_val(13, 10, data["motivo"] if data["motivo"] else "-")
    
    # Observación
    set_val(12, 11, "Observación")
    set_val(13, 11, data["obs"] if data["obs"] else "-")


def generar_excel_skus(skus: list[str]) -> bytes:
    if not skus:
        raise ValueError("Lista de SKUs vac\u00eda")

    from src.services.maquila_service import build_maquila_families
    with engine.connect() as conn:
        df_lookup = pd.read_sql(text("SELECT codigo_femaco, sku, nombre_producto, stock_act, total_4_sem_verificado FROM planificacion_sop"), conn)
    df_lookup["codigo_femaco"] = (
        df_lookup["codigo_femaco"].fillna("").astype(str).str.strip().str.upper()
    )
    lookup_dict = df_lookup.set_index("codigo_femaco")[["sku", "nombre_producto", "stock_act", "total_4_sem_verificado"]].to_dict("index")
    for k, v in lookup_dict.items():
        v["ritmo_mensual"] = v.pop("total_4_sem_verificado", 0)
    
    with engine.connect() as conn:
        familias_map = build_maquila_families(conn, lookup_dict)

    datos_skus = []
    faltantes = []
    for s in skus:
        d = get_sku_export_data(s, familias_map=familias_map)
        if not d:
            faltantes.append(s)
        else:
            datos_skus.append(d)
            
    if faltantes:
        raise ValueError(f"Los siguientes SKUs no se encontraron o no tienen datos: {', '.join(faltantes)}")

    # Agrupar por categoria
    agrupados = {}
    for d in datos_skus:
        cat = d["categoria"]
        if cat not in agrupados:
            agrupados[cat] = []
        agrupados[cat].append(d)
        
    # Ordenar interno por sku
    for cat in agrupados:
        agrupados[cat].sort(key=lambda x: str(x["sku"]))
        
    base_dir = Path(__file__).resolve().parent.parent
    template_path = base_dir / "templates" / "RRFF_AS_PLANTILLA_APP.xlsx"
    if not template_path.exists():
        raise FileNotFoundError(f"Plantilla no encontrada en {template_path}")

    wb = openpyxl.load_workbook(template_path)
    ws_plantilla = wb["PLANTILLA"]
    TAMANO_BLOQUE = 16

    # Ordenar categorias alfabeticamente para las hojas
    categorias_ordenadas = sorted(agrupados.keys())
    
    for cat in categorias_ordenadas:
        sheet_name = CATEGORY_SHEET_MAP.get(cat, cat.title()[:31])
        ws_cat = wb.copy_worksheet(ws_plantilla)
        ws_cat.title = sheet_name
        
        # Poblar bloques en la hoja de la categoria
        for i, data in enumerate(agrupados[cat]):
            fila_inicio = 1 + (i * TAMANO_BLOQUE)
            
            if i > 0:
                copiar_bloque(ws_cat, fila_origen_inicio=1, fila_origen_fin=TAMANO_BLOQUE, fila_destino_inicio=fila_inicio)
                
            poblar_bloque(ws_cat, fila_inicio, data)
            
    # Eliminar plantillas originales
    if "PLANTILLA" in wb.sheetnames:
        del wb["PLANTILLA"]
    if "MAPEO_DEV" in wb.sheetnames:
        del wb["MAPEO_DEV"]

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def generate_excel_for_sku(sku: str) -> bytes:
    """Wrapper para mantener compatibilidad con el Hito 1."""
    return generar_excel_skus([sku])

def generar_excel_todos_skus() -> bytes:
    """Genera un archivo Excel único conteniendo todos los SKUs de planificacion_sop en orden alfanumérico."""
    with engine.connect() as conn:
        df = pd.read_sql(text("SELECT sku FROM planificacion_sop ORDER BY sku ASC"), conn)
        
    skus_totales = df["sku"].dropna().tolist()
    if not skus_totales:
        raise ValueError("No se encontraron SKUs en planificacion_sop")
        
    return generar_excel_skus(skus_totales)
